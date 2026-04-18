"""Shared helpers for NVDA workspace Excel build scripts."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = REPO_ROOT / "data"
WORKBOOKS_DIR = REPO_ROOT / "workbooks"

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
SUBHEADER_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
DEFAULT_FONT = Font(name="Calibri", size=10, color="111827")
MUTED_FONT = Font(name="Calibri", size=9, italic=True, color="6B7280")
THIN = Side(border_style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


PERIOD_RE = re.compile(r"FY(\d{2})(?:Q(\d))?$", re.IGNORECASE)


@dataclass(order=True)
class Period:
    """Sortable fiscal-period key. Annual periods sort after the Q4 of that year."""

    fy: int
    q: int  # 1-4 for quarters, 5 for annual (so FY26 annual > FY26Q4)

    @classmethod
    def parse(cls, text: str) -> "Period":
        m = PERIOD_RE.match(text.strip())
        if not m:
            raise ValueError(f"bad period: {text!r}")
        fy = int(m.group(1))
        q = int(m.group(2)) if m.group(2) else 5
        return cls(fy=fy, q=q)

    @property
    def is_annual(self) -> bool:
        return self.q == 5

    @property
    def label(self) -> str:
        return f"FY{self.fy:02d}" if self.is_annual else f"FY{self.fy:02d}Q{self.q}"


def discover_yaml(dir_path: Path) -> list[Path]:
    """Return all .yaml files under a directory, recursively, sorted."""
    if not dir_path.exists():
        return []
    return sorted(dir_path.rglob("*.yaml"))


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def style_subheader(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = SUBHEADER_FILL
        cell.font = SUBHEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BORDER


def autosize(ws: Worksheet, min_width: int = 10, max_width: int = 40) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_width
        for cell in col:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            if length > width:
                width = length
        ws.column_dimensions[letter].width = min(max_width, width + 2)


def write_value(ws: Worksheet, row: int, col: int, value: Any, *, fmt: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = DEFAULT_FONT
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    if isinstance(value, (int, float)):
        cell.alignment = Alignment(horizontal="right")
    else:
        cell.alignment = Alignment(horizontal="left")


def muted(ws: Worksheet, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = MUTED_FONT
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def ensure_workbooks_dir() -> None:
    WORKBOOKS_DIR.mkdir(parents=True, exist_ok=True)


def new_workbook() -> Workbook:
    wb = Workbook()
    # strip the default "Sheet"
    default = wb.active
    wb.remove(default)
    return wb


def add_sheet(wb: Workbook, name: str) -> Worksheet:
    ws = wb.create_sheet(title=name)
    ws.sheet_view.showGridLines = False
    return ws


def get_scalar(node: Any) -> Any:
    """Extract the numeric value from either a {value:, source:} dict or a scalar."""
    if isinstance(node, dict):
        return node.get("value")
    return node


def get_source(node: Any) -> str:
    if isinstance(node, dict):
        return node.get("source") or ""
    return ""


def flatten_periods(files: Iterable[Path]) -> list[tuple[Period, Path]]:
    out: list[tuple[Period, Path]] = []
    for f in files:
        stem = f.stem  # e.g. FY25 or FY26Q1
        try:
            out.append((Period.parse(stem), f))
        except ValueError:
            continue
    out.sort(key=lambda t: t[0])
    return out
