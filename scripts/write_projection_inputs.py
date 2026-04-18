"""Write values from data/assumptions/base.yaml into named input cells of
workbooks/NVDA_Projections.xlsx.

Only touches cells addressable via named ranges registered on the allowed
sheets. Refuses to write anywhere else. This is the guard that prevents
corrupting the hand-built formula model.

Allowed sheets (per CLAUDE.md):
  Assumptions_BottomUp, Assumptions_TopDown, Supply_Model, Margins, Opex_Below

Name prefix: all assumption named ranges begin with `nv.`. The flattened
YAML key is appended (dots preserved).
"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Iterable

import yaml
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

from lib.excel_utils import DATA_DIR, WORKBOOKS_DIR

WORKBOOK = WORKBOOKS_DIR / "NVDA_Projections.xlsx"
ASSUMPTIONS = DATA_DIR / "assumptions" / "base.yaml"

ALLOWED_SHEETS = {
    "Assumptions_BottomUp",
    "Assumptions_TopDown",
    "Supply_Model",
    "Margins",
    "Opex_Below",
}
NAME_PREFIX = "nv."


def flatten(doc: Any, prefix: str = "") -> Iterable[tuple[str, Any]]:
    """Flatten a nested dict to dotted keys."""
    if isinstance(doc, dict):
        for k, v in doc.items():
            yield from flatten(v, f"{prefix}{k}." if prefix else f"{k}.")
    else:
        # leaf — trim trailing dot from prefix
        yield prefix.rstrip("."), doc


def resolve_named_range(wb, name: str) -> tuple[str, str] | None:
    """Return (sheet_name, cell_ref) for a defined name, or None if not found
    or not resolvable to a single cell."""
    defn: DefinedName | None = wb.defined_names.get(name)
    if defn is None:
        return None
    destinations = list(defn.destinations)
    if len(destinations) != 1:
        return None
    sheet_name, cell_ref = destinations[0]
    # cell_ref might be like '$B$5' or a range '$B$5:$C$5'
    if ":" in cell_ref:
        return None
    return sheet_name, cell_ref.replace("$", "")


def main() -> int:
    if not WORKBOOK.exists():
        print(f"error: {WORKBOOK} not found — build the projection model first", file=sys.stderr)
        return 1
    if not ASSUMPTIONS.exists():
        print(f"error: {ASSUMPTIONS} not found", file=sys.stderr)
        return 1

    with open(ASSUMPTIONS) as f:
        doc = yaml.safe_load(f) or {}

    wb = load_workbook(WORKBOOK, keep_vba=False)

    wrote = 0
    skipped_missing = []
    skipped_disallowed = []
    skipped_bad_range = []
    skipped_null = 0

    for key, value in flatten(doc):
        # Skip null values: keep the cell's current content rather than wiping.
        if value is None:
            skipped_null += 1
            continue
        name = f"{NAME_PREFIX}{key}"
        resolved = resolve_named_range(wb, name)
        if resolved is None:
            if wb.defined_names.get(name) is None:
                skipped_missing.append(name)
            else:
                skipped_bad_range.append(name)
            continue
        sheet_name, cell_ref = resolved
        if sheet_name not in ALLOWED_SHEETS:
            skipped_disallowed.append((name, sheet_name))
            continue
        ws = wb[sheet_name]
        ws[cell_ref] = value
        wrote += 1

    wb.save(WORKBOOK)

    print(f"wrote {wrote} value(s) into {WORKBOOK}")
    if skipped_null:
        print(f"  {skipped_null} key(s) null in YAML — workbook cells left unchanged")
    if skipped_missing:
        print(f"  {len(skipped_missing)} key(s) with no matching named range (first 5): "
              + ", ".join(skipped_missing[:5]))
    if skipped_bad_range:
        print(f"  {len(skipped_bad_range)} named range(s) not resolvable to a single cell: "
              + ", ".join(skipped_bad_range[:5]))
    if skipped_disallowed:
        names = [f"{n}@{s}" for n, s in skipped_disallowed[:5]]
        print(f"  {len(skipped_disallowed)} named range(s) on disallowed sheet (refused): "
              + ", ".join(names))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
