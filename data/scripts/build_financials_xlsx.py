"""Build workbooks/NVDA_Financials.xlsx from data/financials/**/*.yaml.

Auto-discovers every annual and quarterly YAML, sorts chronologically, and
emits a lean P&L / Segments / Cash & Debt / Cash Flow / Sources workbook.
New periods drop in without script changes.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import yaml
from openpyxl.styles import Alignment, Font

from lib.excel_utils import (
    DATA_DIR,
    WORKBOOKS_DIR,
    Period,
    add_sheet,
    autosize,
    discover_yaml,
    ensure_workbooks_dir,
    flatten_periods,
    get_scalar,
    get_source,
    muted,
    new_workbook,
    style_header,
    style_subheader,
    write_value,
)

OUTPUT = WORKBOOKS_DIR / "NVDA_Financials.xlsx"

# Lean v1 scope — rows emitted on each sheet. Each entry: (label, yaml path, format)
PNL_ROWS = [
    ("Revenue ($m)", ["income_statement", "revenue_total"], "#,##0"),
    ("  Data Center", ["income_statement", "revenue_by_segment", "data_center"], "#,##0"),
    ("  Gaming", ["income_statement", "revenue_by_segment", "gaming"], "#,##0"),
    ("  Pro Viz", ["income_statement", "revenue_by_segment", "pro_viz"], "#,##0"),
    ("  Auto", ["income_statement", "revenue_by_segment", "auto"], "#,##0"),
    ("  OEM", ["income_statement", "revenue_by_segment", "oem"], "#,##0"),
    ("Gross profit (GAAP, $m)", ["income_statement", "gross_profit_gaap"], "#,##0"),
    ("GM% (GAAP)", ["income_statement", "gross_margin_gaap_pct"], "0.0%"),
    ("GM% (non-GAAP)", ["income_statement", "gross_margin_nongaap_pct"], "0.0%"),
    ("Opex — R&D ($m)", ["income_statement", "opex_rd"], "#,##0"),
    ("Opex — SG&A ($m)", ["income_statement", "opex_sga"], "#,##0"),
    ("Operating income (GAAP, $m)", ["income_statement", "operating_income_gaap"], "#,##0"),
    ("Net income (GAAP, $m)", ["income_statement", "net_income_gaap"], "#,##0"),
    ("EPS diluted (GAAP, $)", ["income_statement", "eps_diluted_gaap"], "0.00"),
    ("EPS diluted (non-GAAP, $)", ["income_statement", "eps_diluted_nongaap"], "0.00"),
]

CASH_DEBT_ROWS = [
    ("Cash + marketable securities ($m)", ["balance_sheet", "cash_and_marketable_securities"], "#,##0"),
    ("Total debt ($m)", ["balance_sheet", "total_debt"], "#,##0"),
    ("Inventory ($m)", ["balance_sheet", "inventory"], "#,##0"),
    ("Purchase commitments ($m)", ["balance_sheet", "purchase_commitments"], "#,##0"),
]

CF_ROWS = [
    ("Operating cash flow ($m)", ["cash_flow", "operating_cash_flow"], "#,##0"),
    ("Capex ($m)", ["cash_flow", "capex"], "#,##0"),
    ("Buybacks ($m)", ["cash_flow", "buybacks"], "#,##0"),
    ("Dividends ($m)", ["cash_flow", "dividends"], "#,##0"),
]

SEGMENT_ROWS = [
    ("Data Center ($m)", ["income_statement", "revenue_by_segment", "data_center"], "#,##0"),
    ("Gaming ($m)", ["income_statement", "revenue_by_segment", "gaming"], "#,##0"),
    ("Pro Viz ($m)", ["income_statement", "revenue_by_segment", "pro_viz"], "#,##0"),
    ("Auto ($m)", ["income_statement", "revenue_by_segment", "auto"], "#,##0"),
    ("OEM ($m)", ["income_statement", "revenue_by_segment", "oem"], "#,##0"),
]


def load_yaml(path: Path) -> dict:
    with open(path) as f:
        return yaml.safe_load(f) or {}


def dig(doc: dict, path: list[str]):
    node = doc
    for key in path:
        if not isinstance(node, dict):
            return None
        node = node.get(key)
        if node is None:
            return None
    return node


def write_rows(ws, periods, docs, rows, start_row=3):
    """Write a row-set (label + per-period values) to a worksheet."""
    # header
    ws.cell(row=start_row, column=1, value="Line item")
    for i, (p, _) in enumerate(periods, start=2):
        ws.cell(row=start_row, column=i, value=p.label)
    style_header(ws, start_row, ncols=1 + len(periods))

    for r, (label, yaml_path, fmt) in enumerate(rows, start=start_row + 1):
        write_value(ws, r, 1, label)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        for i, (_, yaml_file) in enumerate(periods, start=2):
            node = dig(docs[yaml_file], yaml_path)
            v = get_scalar(node)
            if v is None:
                write_value(ws, r, i, "")
            elif "pct" in "/".join(yaml_path) and isinstance(v, (int, float)) and v > 1:
                # percentages stored as 74.2 → write as 0.742
                write_value(ws, r, i, v / 100.0, fmt=fmt)
            else:
                write_value(ws, r, i, v, fmt=fmt)


def build_cover(wb, n_quarterly, n_annual, periods):
    ws = add_sheet(wb, "Cover")
    ws["A1"] = "NVDA Financials — generated workbook"
    ws["A1"].font = Font(name="Calibri", bold=True, size=16, color="111827")
    ws["A3"] = f"Last build: {dt.datetime.now():%Y-%m-%d %H:%M}"
    ws["A4"] = f"Quarterly YAMLs: {n_quarterly}    Annual YAMLs: {n_annual}"
    ws["A6"] = "Periods covered:"
    ws["A7"] = ", ".join(p.label for p, _ in periods)
    ws["A9"] = "Rebuild with:  python scripts/build_financials_xlsx.py"
    ws["A10"] = "Source of truth: data/financials/**/*.yaml — do NOT hand-edit this workbook."
    ws.column_dimensions["A"].width = 100


def build_sheet(wb, name, periods, docs, rows):
    ws = add_sheet(wb, name)
    ws["A1"] = f"NVDA — {name}"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color="111827")
    write_rows(ws, periods, docs, rows, start_row=3)
    autosize(ws, min_width=14, max_width=30)


def build_sources(wb, periods, docs, row_defs):
    ws = add_sheet(wb, "Sources")
    ws["A1"] = "Source citations per line item"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color="111827")

    ws.cell(row=3, column=1, value="Line item")
    ws.cell(row=3, column=2, value="Period")
    ws.cell(row=3, column=3, value="Source")
    ws.cell(row=3, column=4, value="URL")
    style_header(ws, 3, ncols=4)

    r = 4
    for label, yaml_path, _ in row_defs:
        for p, yaml_file in periods:
            node = dig(docs[yaml_file], yaml_path)
            src = get_source(node)
            if not src:
                continue
            url = docs[yaml_file].get("source_urls", {})
            url_str = ""
            if isinstance(url, dict):
                # pick a plausible URL key
                for k in ("press_release", "ten_q", "ten_k", "cfo_commentary", "transcript"):
                    if url.get(k):
                        url_str = url[k]
                        break
            write_value(ws, r, 1, label)
            write_value(ws, r, 2, p.label)
            write_value(ws, r, 3, src)
            muted(ws, r, 4, url_str)
            r += 1
    autosize(ws, min_width=16, max_width=60)


def main() -> None:
    ensure_workbooks_dir()

    quarterly_files = discover_yaml(DATA_DIR / "financials" / "quarterly")
    annual_files = discover_yaml(DATA_DIR / "financials" / "annual")

    quarterly = flatten_periods(quarterly_files)
    annual = flatten_periods(annual_files)

    docs = {p: load_yaml(p) for p in quarterly_files + annual_files}

    wb = new_workbook()
    build_cover(wb, len(quarterly), len(annual), quarterly + annual)

    if quarterly:
        build_sheet(wb, "P&L (Quarterly)", quarterly, docs, PNL_ROWS)
        build_sheet(wb, "Segments (Quarterly)", quarterly, docs, SEGMENT_ROWS)
        build_sheet(wb, "Cash & Debt (Quarterly)", quarterly, docs, CASH_DEBT_ROWS)
        build_sheet(wb, "Cash Flow (Quarterly)", quarterly, docs, CF_ROWS)

    if annual:
        build_sheet(wb, "P&L (Annual)", annual, docs, PNL_ROWS)
        build_sheet(wb, "Segments (Annual)", annual, docs, SEGMENT_ROWS)
        build_sheet(wb, "Cash & Debt (Annual)", annual, docs, CASH_DEBT_ROWS)
        build_sheet(wb, "Cash Flow (Annual)", annual, docs, CF_ROWS)

    build_sources(wb, quarterly + annual, docs, PNL_ROWS + SEGMENT_ROWS + CASH_DEBT_ROWS + CF_ROWS)

    wb.save(OUTPUT)
    print(f"wrote {OUTPUT}  ({len(quarterly)} quarterly + {len(annual)} annual)")


if __name__ == "__main__":
    main()
