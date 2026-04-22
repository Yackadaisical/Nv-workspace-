"""Build the hand-built projection workbook template.

This script is run ONCE (or after schema changes) to produce
workbooks/NVDA_Projections.xlsx with:
  - Dashboard, Scenarios, Assumptions_BottomUp, Assumptions_TopDown,
    Supply_Model, Margins, Opex_Below, Reconciliation, Change_Log_Mirror
    sheets
  - All named ranges (prefix `nv.`) wired up so write_projection_inputs.py
    can push values from data/assumptions/base.yaml
  - Live formulas on Dashboard and Reconciliation sheets that reference the
    input cells

After this runs, the user is free to hand-edit formulas/formatting on the
workbook. Re-running this script will OVERWRITE the workbook, so don't run
it twice without backing up.

This is the v1 skeleton. The model can be expanded by the user as needed —
additional SKUs, customers, quarters — just add the same pattern of named
ranges and they become writable by write_projection_inputs.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

WORKBOOK = Path(__file__).resolve().parents[1] / "workbooks" / "NVDA_Projections.xlsx"

FWD_QUARTERS = ["fy27q1", "fy27q2", "fy27q3", "fy27q4",
                "fy28q1", "fy28q2", "fy28q3", "fy28q4"]
FWD_YEARS = ["fy27", "fy28", "fy29"]

CUSTOMERS = ["msft", "meta", "goog", "amzn", "orcl", "tier2", "sovereign", "enterprise"]
SKUS = ["gb200"]  # v1: one SKU. User can extend by adding named ranges.

THIN = Side(border_style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FEF3C7")   # yellow: input
FORMULA_FILL = PatternFill("solid", fgColor="DBEAFE") # blue: formula
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
BODY_FONT = Font(name="Calibri", size=10, color="111827")


def add_name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    """Register a workbook-scoped defined name pointing at one cell."""
    ref = f"'{sheet}'!${cell[0]}${cell[1:]}"
    wb.defined_names[name] = DefinedName(name=name, attr_text=ref)


def style_input(ws: Worksheet, cell: str) -> None:
    c = ws[cell]
    c.fill = INPUT_FILL
    c.font = BODY_FONT
    c.border = BORDER
    c.alignment = Alignment(horizontal="right")
    c.number_format = "#,##0.00"


def style_formula(ws: Worksheet, cell: str) -> None:
    c = ws[cell]
    c.fill = FORMULA_FILL
    c.font = BODY_FONT
    c.border = BORDER
    c.alignment = Alignment(horizontal="right")
    c.number_format = "#,##0"


def style_header(ws: Worksheet, cell: str, value) -> None:
    c = ws[cell]
    c.value = value
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BORDER


def style_label(ws: Worksheet, cell: str, value) -> None:
    c = ws[cell]
    c.value = value
    c.font = LABEL_FONT
    c.alignment = Alignment(horizontal="left")


def cell_ref(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def build_scenarios(wb: Workbook) -> None:
    ws = wb.create_sheet("Scenarios")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Scenario toggle")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    style_label(ws, "A3", "Active scenario (type: Base / Bear / Bull)")
    ws["B3"] = "Base"
    style_input(ws, "B3")
    ws["B3"].number_format = "@"
    add_name(wb, "nv.scenario_selector", "Scenarios", "B3")
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 18


def build_bottom_up(wb: Workbook) -> None:
    ws = wb.create_sheet("Assumptions_BottomUp")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Bottom-up demand (units × ASP by customer × SKU × quarter)")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    style_label(ws, "A2", "Yellow = input. Revenue columns below are formula = units × ASP.")
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="6B7280")

    row = 4
    for cust in CUSTOMERS:
        for sku in SKUS:
            # section header
            style_header(ws, cell_ref(1, row), f"{cust.upper()} × {sku.upper()}")
            for i, q in enumerate(FWD_QUARTERS, start=2):
                style_header(ws, cell_ref(i, row), q)
            row += 1

            # units row
            style_label(ws, cell_ref(1, row), "Units")
            for i, q in enumerate(FWD_QUARTERS, start=2):
                c = cell_ref(i, row)
                style_input(ws, c)
                add_name(wb, f"nv.bu.{cust}.{sku}.units.{q}", "Assumptions_BottomUp", c)
            row += 1

            # asp row
            style_label(ws, cell_ref(1, row), "ASP ($)")
            for i, q in enumerate(FWD_QUARTERS, start=2):
                c = cell_ref(i, row)
                style_input(ws, c)
                add_name(wb, f"nv.bu.{cust}.{sku}.asp.{q}", "Assumptions_BottomUp", c)
            row += 1

            # revenue row (formula = units × asp / 1e6, in $m)
            style_label(ws, cell_ref(1, row), "Revenue ($m)")
            for i, q in enumerate(FWD_QUARTERS, start=2):
                c = cell_ref(i, row)
                units = cell_ref(i, row - 2)
                asp = cell_ref(i, row - 1)
                ws[c] = f"=IFERROR({units}*{asp}/1000000, 0)"
                style_formula(ws, c)
            row += 2  # blank separator

    # Totals block
    style_header(ws, cell_ref(1, row), "TOTAL BU REVENUE ($m)")
    for i, q in enumerate(FWD_QUARTERS, start=2):
        style_header(ws, cell_ref(i, row), q)
    row += 1
    style_label(ws, cell_ref(1, row), "Total units (all customers × all SKUs)")
    for i, q in enumerate(FWD_QUARTERS, start=2):
        # Sum every 4th row starting from first units row in each customer block
        # Simpler: sum the named ranges for that quarter
        refs = [f"nv.bu.{c}.{s}.units.{q}" for c in CUSTOMERS for s in SKUS]
        ws[cell_ref(i, row)] = "=" + "+".join(refs)
        style_formula(ws, cell_ref(i, row))
    row += 1
    style_label(ws, cell_ref(1, row), "Total revenue ($m)")
    for i, q in enumerate(FWD_QUARTERS, start=2):
        refs = [f"nv.bu.{c}.{s}.units.{q}*nv.bu.{c}.{s}.asp.{q}" for c in CUSTOMERS for s in SKUS]
        ws[cell_ref(i, row)] = "=(" + "+".join(refs) + ")/1000000"
        style_formula(ws, cell_ref(i, row))
        add_name(wb, f"nv.bu.total_rev.{q}", "Assumptions_BottomUp", cell_ref(i, row))

    # column widths
    ws.column_dimensions["A"].width = 36
    for i in range(2, 2 + len(FWD_QUARTERS)):
        ws.column_dimensions[get_column_letter(i)].width = 12


def build_top_down(wb: Workbook) -> None:
    ws = wb.create_sheet("Assumptions_TopDown")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Top-down: accelerator TAM × NVDA share")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)

    style_header(ws, "A3", "Metric")
    for i, y in enumerate(FWD_YEARS, start=2):
        style_header(ws, cell_ref(i, 3), y)

    style_label(ws, "A4", "Accelerator TAM ($m)")
    for i, y in enumerate(FWD_YEARS, start=2):
        c = cell_ref(i, 4)
        style_input(ws, c)
        add_name(wb, f"nv.td.accelerator_tam.{y}", "Assumptions_TopDown", c)

    style_label(ws, "A5", "NVDA share (0..1)")
    for i, y in enumerate(FWD_YEARS, start=2):
        c = cell_ref(i, 5)
        style_input(ws, c)
        ws[c].number_format = "0.0%"
        add_name(wb, f"nv.td.nvda_share.{y}", "Assumptions_TopDown", c)

    style_label(ws, "A6", "Implied NVDA revenue ($m)")
    for i, y in enumerate(FWD_YEARS, start=2):
        tam = cell_ref(i, 4)
        share = cell_ref(i, 5)
        c = cell_ref(i, 6)
        ws[c] = f"=IFERROR({tam}*{share}, 0)"
        style_formula(ws, c)
        add_name(wb, f"nv.td.implied_rev.{y}", "Assumptions_TopDown", c)

    ws.column_dimensions["A"].width = 34
    for i in range(2, 2 + len(FWD_YEARS)):
        ws.column_dimensions[get_column_letter(i)].width = 14


def build_supply(wb: Workbook) -> None:
    ws = wb.create_sheet("Supply_Model")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Supply model (annual)")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)

    style_header(ws, "A3", "Metric")
    for i, y in enumerate(FWD_YEARS, start=2):
        style_header(ws, cell_ref(i, 3), y)

    rows = [
        ("CoWoS wafers available", "supply.cowos_wafers", "#,##0"),
        ("HBM allocation (GB)",    "supply.hbm_allocation_gb", "#,##0"),
        ("Packaging yield (0..1)", "supply.packaging_yield_pct", "0.0%"),
    ]
    for r_idx, (label, key_suffix, fmt) in enumerate(rows, start=4):
        style_label(ws, cell_ref(1, r_idx), label)
        for i, y in enumerate(FWD_YEARS, start=2):
            c = cell_ref(i, r_idx)
            style_input(ws, c)
            ws[c].number_format = fmt
            add_name(wb, f"nv.{key_suffix}.{y}", "Supply_Model", c)

    ws.column_dimensions["A"].width = 34
    for i in range(2, 2 + len(FWD_YEARS)):
        ws.column_dimensions[get_column_letter(i)].width = 14


def build_margins(wb: Workbook) -> None:
    ws = wb.create_sheet("Margins")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Margin drivers")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)

    style_header(ws, "A3", "Driver")
    for i, q in enumerate(FWD_QUARTERS, start=2):
        style_header(ws, cell_ref(i, 3), q)

    style_label(ws, "A4", "HBM cost per GB ($)")
    for i, q in enumerate(FWD_QUARTERS, start=2):
        c = cell_ref(i, 4)
        style_input(ws, c)
        add_name(wb, f"nv.margin.hbm_cost_per_gb.{q}", "Margins", c)

    style_label(ws, "A5", "Yield uplift (%)")
    for i, q in enumerate(FWD_QUARTERS, start=2):
        c = cell_ref(i, 5)
        style_input(ws, c)
        ws[c].number_format = "0.0%"
        add_name(wb, f"nv.margin.yield_uplift_pct.{q}", "Margins", c)

    ws.column_dimensions["A"].width = 34
    for i in range(2, 2 + len(FWD_QUARTERS)):
        ws.column_dimensions[get_column_letter(i)].width = 12


def build_opex(wb: Workbook) -> None:
    ws = wb.create_sheet("Opex_Below")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Opex & below-the-line (annual)")
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)

    style_header(ws, "A3", "Driver")
    for i, y in enumerate(FWD_YEARS, start=2):
        style_header(ws, cell_ref(i, 3), y)

    rows = [
        ("R&D YoY growth (0..1)", "opex.rd_growth_yoy_pct", "0.0%"),
        ("SG&A YoY growth (0..1)", "opex.sga_growth_yoy_pct", "0.0%"),
        ("Tax rate (0..1)", "opex.tax_rate_pct", "0.0%"),
        ("Buybacks ($m)", "opex.buyback_usd", "#,##0"),
    ]
    for r_idx, (label, key_suffix, fmt) in enumerate(rows, start=4):
        style_label(ws, cell_ref(1, r_idx), label)
        for i, y in enumerate(FWD_YEARS, start=2):
            c = cell_ref(i, r_idx)
            style_input(ws, c)
            ws[c].number_format = fmt
            add_name(wb, f"nv.{key_suffix}.{y}", "Opex_Below", c)

    ws.column_dimensions["A"].width = 34
    for i in range(2, 2 + len(FWD_YEARS)):
        ws.column_dimensions[get_column_letter(i)].width = 14


def build_dashboard(wb: Workbook) -> None:
    ws = wb.create_sheet("Dashboard", 0)  # make it first
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "NVDA forecast dashboard")
    ws["A1"].font = Font(name="Calibri", size=16, bold=True)
    style_label(ws, "A2", "Active scenario:")
    ws["B2"] = "=nv.scenario_selector"
    ws["B2"].font = LABEL_FONT

    # Forward annual revenue summary (sum of 4 quarters of bottom-up)
    style_label(ws, "A4", "Annual revenue ($m) — bottom-up")
    style_header(ws, "A5", "FY")
    style_header(ws, "B5", "Bottom-up rev ($m)")
    style_header(ws, "C5", "Top-down implied ($m)")
    style_header(ws, "D5", "Reconciliation gap")

    for i, y in enumerate(FWD_YEARS, start=6):
        style_label(ws, cell_ref(1, i), y)
        # Bottom-up = sum of 4 forward quarters if available
        quarters_of_year = [q for q in FWD_QUARTERS if q.startswith(y)]
        if quarters_of_year:
            bu_refs = [f"nv.bu.total_rev.{q}" for q in quarters_of_year]
            ws[cell_ref(2, i)] = "=IFERROR(" + "+".join(bu_refs) + ", 0)"
        else:
            ws[cell_ref(2, i)] = 0
        style_formula(ws, cell_ref(2, i))
        # Top-down implied
        ws[cell_ref(3, i)] = f"=IFERROR(nv.td.implied_rev.{y}, 0)"
        style_formula(ws, cell_ref(3, i))
        # Gap (BU - TD)
        ws[cell_ref(4, i)] = f"={cell_ref(2, i)}-{cell_ref(3, i)}"
        style_formula(ws, cell_ref(4, i))

    ws["A11"] = "Note: populate input cells on Assumptions_BottomUp, Assumptions_TopDown, Supply_Model, Margins, and Opex_Below — or run scripts/write_projection_inputs.py to push from data/assumptions/base.yaml."
    ws["A11"].font = Font(name="Calibri", size=9, italic=True, color="6B7280")
    ws["A11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A11:D13")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22


def build_reconciliation(wb: Workbook) -> None:
    ws = wb.create_sheet("Reconciliation")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Model vs Street vs Company guide (placeholder — fill in as consensus is gathered)")
    ws["A1"].font = Font(name="Calibri", size=12, bold=True)

    style_header(ws, "A3", "Period")
    style_header(ws, "B3", "Model ($m)")
    style_header(ws, "C3", "Street consensus ($m)")
    style_header(ws, "D3", "Company guide ($m)")
    style_header(ws, "E3", "Notes")

    for i, q in enumerate(FWD_QUARTERS, start=4):
        style_label(ws, cell_ref(1, i), q)
        ws[cell_ref(2, i)] = f"=IFERROR(nv.bu.total_rev.{q}, 0)"
        style_formula(ws, cell_ref(2, i))
        # user fills C, D, E

    ws.column_dimensions["A"].width = 10
    for col, w in [("B", 16), ("C", 22), ("D", 22), ("E", 40)]:
        ws.column_dimensions[col].width = w


def build_changelog_mirror(wb: Workbook) -> None:
    ws = wb.create_sheet("Change_Log_Mirror")
    ws.sheet_view.showGridLines = False
    style_label(ws, "A1", "Change_Log_Mirror — reflects last 20 entries of views/changelog.md")
    ws["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws["A2"] = "This sheet is intended to be refreshed alongside projection updates. For v1, mirror the file manually or extend scripts to sync."
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="6B7280")
    ws.column_dimensions["A"].width = 120


def main() -> int:
    wb = Workbook()
    # remove default Sheet
    default = wb.active
    wb.remove(default)

    build_scenarios(wb)
    build_bottom_up(wb)
    build_top_down(wb)
    build_supply(wb)
    build_margins(wb)
    build_opex(wb)
    build_reconciliation(wb)
    build_changelog_mirror(wb)
    # Dashboard last so it references already-defined named ranges; placed at index 0 inside the function
    build_dashboard(wb)

    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)
    print(f"wrote {WORKBOOK} with {len(wb.defined_names)} named range(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
